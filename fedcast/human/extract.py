"""Text extraction for uploaded documents. Code only: the verbatim text is the evidence.

An LLM summary is never stored as evidence; summarising is an analyst step that must quote
this text. Supported: .txt, .md, .pdf, .docx.
"""

from __future__ import annotations

import io
from pathlib import PurePath

SUPPORTED = (".txt", ".md", ".pdf", ".docx")


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext in (".txt", ".md"):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError(f"unsupported file type {ext or '(none)'}; use one of {', '.join(SUPPORTED)}")
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < 20:
        raise ValueError("no readable text found in the file (a scanned PDF needs OCR first)")
    return text
