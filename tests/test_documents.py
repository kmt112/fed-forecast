from datetime import date

import pytest

from fedcast import dag
from fedcast.human import documents as ds
from fedcast.snapshot import Item
from tests.test_journey_ui import entry_for, full_snapshot


def doc(id_="H5-001", **kw):
    base = dict(id=id_, title="Bank lending standards note", author="t", added_on=date(2026, 10, 2),
                source="own analysis", relevance="related",
                text="Senior loan officer survey shows standards tightening for the third quarter running.")
    return ds.Document(**{**base, **kw})


def test_roundtrip_through_markdown(tmp_path):
    ds.save(tmp_path, doc())
    loaded = ds.load_documents(tmp_path)
    assert loaded == [doc()]
    assert ds.next_id(loaded) == "H5-002"


def test_cannot_overwrite_by_accident(tmp_path):
    ds.save(tmp_path, doc())
    with pytest.raises(FileExistsError):
        ds.save(tmp_path, doc(title="Different title"))


def test_free_form_relevance_is_rejected():
    with pytest.raises(Exception):
        doc(relevance="kind of")


def test_documents_appear_in_the_graph_as_frozen_evidence_feeding_only_the_analysts():
    snap = full_snapshot()
    snap.add(Item.make("human.doc.H5-001", "human_document", "human/documents/H5-001.md", doc().model_dump(mode="json")))
    g = dag.build(entry_for(snap, []), snap)
    node = next(n for n in g["nodes"] if n["id"] == "documents")
    assert node["status"] == "idle" and node["detail"]["rows"][0][0] == "H5-001"
    targets = {e["to"] for e in g["edges"] if e["from"] == "documents"}
    assert targets == {"model.llm"}
    assert all(e["status"] == "planned" for e in g["edges"] if e["from"] == "documents")


def test_extract_text_from_docx_and_txt():
    import io

    import docx

    from fedcast.human.extract import extract_text

    d = docx.Document()
    d.add_paragraph("Lending standards tightened for a third quarter.")
    buf = io.BytesIO(); d.save(buf)
    assert "tightened for a third quarter" in extract_text("note.docx", buf.getvalue())
    assert extract_text("n.txt", "Core PCE ran at 3.3% over the year to July.".encode()) .startswith("Core PCE")
    with pytest.raises(ValueError):
        extract_text("x.exe", b"binary")


def test_upload_goes_through_the_same_document_model(tmp_path, monkeypatch):
    from fedcast.ui import server

    monkeypatch.setattr(server, "DOCS", tmp_path)
    body = "Claims have drifted up for six weeks running, which the Fed will notice.".encode()
    boundary = "XYZ"
    multipart = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"title\"\r\n\r\nClaims note\r\n"
                 f"--{boundary}\r\nContent-Disposition: form-data; name=\"relevance\"\r\n\r\nrelated\r\n"
                 f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"claims.txt\"\r\n"
                 f"Content-Type: text/plain\r\n\r\n").encode() + body + f"\r\n--{boundary}--\r\n".encode()
    fields, filename, data = server._parse_multipart(f"multipart/form-data; boundary={boundary}", multipart)
    assert fields == {"title": "Claims note", "relevance": "related"} and filename == "claims.txt" and data == body
    msg = server.act_upload_document(fields, filename, data)["message"]
    docs = ds.load_documents(tmp_path)
    assert docs[0].id == "H5-001" and docs[0].title == "Claims note" and docs[0].text == body.decode()
    assert (tmp_path / "files" / "H5-001.txt").read_bytes() == body
